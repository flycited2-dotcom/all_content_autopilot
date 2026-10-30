"""Bounded, read-only extraction of owner-supplied documents.

This module returns source data, never instructions or executable content. It
does not evaluate workbook formulas, macros, YAML tags, or document scripts.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 120_000
MAX_ROWS = 200
MAX_COLUMNS = 128
MAX_CELLS = 250_000
MAX_PAGES = 40
MAX_IMAGES = 8
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 20 * 1024 * 1024
SUPPORTED_EXTENSIONS = frozenset({'.pdf', '.docx', '.doc', '.rtf', '.txt', '.csv', '.tsv',
                                  '.yaml', '.yml', '.json', '.xlsx', '.xlsm', '.xls', '.ods'})


class _Output:
    def __init__(self):
        self.lines: list[str] = []
        self.size = 0
        self.rows = 0
        self.images: list[dict] = []
        self.warnings: list[str] = []

    def add(self, text):
        text = str(text)
        self.size += len(text) + 1
        if self.size > MAX_TEXT_CHARS:
            raise ValueError(f'Документ содержит больше {MAX_TEXT_CHARS} символов. Разделите его на несколько файлов')
        if text.strip():
            self.lines.append(text)

    def row(self, index, values):
        values = list(values)
        if len(values) > MAX_COLUMNS:
            raise ValueError(f'В таблице больше {MAX_COLUMNS} столбцов. Разделите файл')
        self.rows += 1
        if self.rows > MAX_ROWS:
            raise ValueError(f'В документе больше {MAX_ROWS} строк таблиц. Разделите файл')
        if any(v is not None and str(v).strip() for v in values):
            self.add(f'Строка {index}: ' + json.dumps(values, ensure_ascii=False, default=str))

    def image(self, data, location):
        from PIL import Image
        if len(self.images) >= MAX_IMAGES:
            raise ValueError(f'В документе больше {MAX_IMAGES} страниц или изображений для распознавания. Разделите файл')
        with Image.open(io.BytesIO(data)) as picture:
            if picture.width * picture.height > 40_000_000:
                raise ValueError('Изображение в документе слишком большое')
            buffer = io.BytesIO()
            picture.convert('RGB').save(buffer, format='JPEG', quality=90)
            rendered = buffer.getvalue()
        if len(rendered) > MAX_IMAGE_BYTES or sum(i['_size'] for i in self.images) + len(rendered) > MAX_TOTAL_IMAGE_BYTES:
            raise ValueError('Изображения документа превышают допустимый размер. Разделите файл')
        self.add(f'{location}: изображение attachment_index={len(self.images)}; требуется распознавание')
        self.images.append({'mime_type': 'image/jpeg', 'data': base64.b64encode(rendered).decode('ascii'),
                            '_size': len(rendered)})


def _decode(data):
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        return data.decode('utf-16')
    if b'\x00' in data:
        raise ValueError('Файл не является текстовым документом')
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        return data.decode('cp1251')


def _xml(data):
    from defusedxml.ElementTree import fromstring
    return fromstring(data)


def _zip(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 2048 or sum(e.file_size for e in entries) > 64 * 1024 * 1024:
        archive.close()
        raise ValueError('Распакованный документ слишком большой. Разделите файл')
    if any(e.flag_bits & 1 or e.file_size > 32 * 1024 * 1024 or
           (e.file_size > 1024 * 1024 and e.file_size > max(1, e.compress_size) * 300)
           for e in entries):
        archive.close()
        raise ValueError('Документ зашифрован или имеет недопустимый размер распаковки')
    if len({e.filename for e in entries}) != len(entries):
        archive.close()
        raise ValueError('В документе обнаружены повторяющиеся внутренние файлы')
    return archive


def _csv(data, out, extension):
    raw = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(raw[:8192], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel_tab if extension == '.tsv' else csv.excel
    for n, row in enumerate(csv.reader(io.StringIO(raw, newline=''), dialect=dialect, strict=True), 1):
        out.row(n, row)


def _structured(data, out, extension):
    raw = _decode(data)
    if len(raw) > MAX_TEXT_CHARS:
        raise ValueError(f'Документ содержит больше {MAX_TEXT_CHARS} символов. Разделите его на несколько файлов')
    if extension == '.json':
        def pairs(entries):
            result = {}
            for k, v in entries:
                if k in result:
                    raise ValueError('В JSON повторяется ключ; исправьте документ, чтобы не потерять данные')
                result[k] = v
            return result
        parsed = json.loads(raw, object_pairs_hook=pairs)
    else:
        import yaml
        # Aliases may expand exponentially even with safe_load. Reject them
        # before construction; supplier tables do not need references.
        tokens = list(yaml.scan(raw))
        if len(tokens) > MAX_CELLS or any(isinstance(t, (yaml.tokens.AliasToken, yaml.tokens.AnchorToken,
                                                       yaml.tokens.TagToken)) for t in tokens):
            raise ValueError('YAML должен содержать обычные данные без ссылок, якорей и специальных тегов')
        class UniqueLoader(yaml.SafeLoader):
            pass
        def unique_mapping(loader, node):
            result = {}
            for key_node, val_node in node.value:
                key = loader.construct_object(key_node, deep=True)
                if key in result:
                    raise ValueError('В YAML повторяется ключ; исправьте документ, чтобы не потерять данные')
                result[key] = loader.construct_object(val_node, deep=True)
            return result
        UniqueLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, unique_mapping)
        parsed = list(yaml.load_all(raw, Loader=UniqueLoader))
    # Keep the exact source spelling, including decimal prices and stock zeros.
    if parsed is None:
        raise ValueError('Документ пуст')
    out.add(raw)


def _xlsx(data, out):
    import openpyxl
    with _zip(data) as archive:
        if 'xl/workbook.xml' not in archive.namelist():
            raise ValueError('Файл не является книгой Excel')
        # The XML is bounded before openpyxl is allowed to parse it.
        for entry in archive.namelist():
            if entry.endswith('.xml'):
                _xml(archive.read(entry))
    values = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        cells = 0
        for sheet, formula_sheet in zip(values.worksheets, formulas.worksheets):
            out.add(f'Лист: {sheet.title}' + (' (скрытый)' if sheet.sheet_state != 'visible' else ''))
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError('Лист Excel превышает предел строк или столбцов. Удалите пустое форматирование или разделите файл')
            for n, (row, formula_row) in enumerate(zip(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ValueError('В Excel слишком много ячеек. Разделите файл')
                extracted = []
                for value, formula in zip(row, formula_row):
                    item = value.value
                    if formula.data_type == 'f' and 'Формулы не вычислялись; использованы сохранённые результаты Excel' not in out.warnings:
                        out.warnings.append('Формулы не вычислялись; использованы сохранённые результаты Excel')
                    if formula.data_type == 'f' and item is None:
                        item = '[формула без сохранённого результата: значение неизвестно]'
                        if 'Формулы без сохранённых результатов не вычислялись' not in out.warnings:
                            out.warnings.append('Формулы без сохранённых результатов не вычислялись')
                    extracted.append(item)
                out.row(n, extracted)
    finally:
        values.close()
        formulas.close()


def _xls(data, out):
    import xlrd
    book = xlrd.open_workbook(file_contents=data, on_demand=True, logfile=io.StringIO())
    try:
        cells = 0
        for sheet in book.sheets():
            if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                raise ValueError('Лист Excel превышает предел строк или столбцов. Разделите файл')
            cells += sheet.nrows * sheet.ncols
            if cells > MAX_CELLS:
                raise ValueError('В Excel слишком много ячеек. Разделите файл')
            out.add('Лист: ' + sheet.name)
            for n in range(sheet.nrows):
                out.row(n + 1, sheet.row_values(n))
    finally:
        book.release_resources()


def _ods(data, out):
    ns = {'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
          'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
          'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
    with _zip(data) as archive:
        root = _xml(archive.read('content.xml'))
    for table in root.findall('.//table:table', ns):
        out.add('Лист: ' + table.get('{%s}name' % ns['table'], ''))
        number = 0
        for row in table.findall('table:table-row', ns):
            repeated_rows = int(row.get('{%s}number-rows-repeated' % ns['table'], '1'))
            values = []
            for cell in row:
                repeated_cells = int(cell.get('{%s}number-columns-repeated' % ns['table'], '1'))
                value = '\n'.join(''.join(p.itertext()) for p in cell.findall('text:p', ns))
                if not value:
                    value = cell.get('{%s}value' % ns['office']) or cell.get('{%s}string-value' % ns['office'])
                if repeated_cells > MAX_COLUMNS or len(values) + repeated_cells > MAX_COLUMNS:
                    raise ValueError('В ODS слишком много столбцов. Разделите файл')
                values.extend([value] * repeated_cells)
            if repeated_rows > MAX_ROWS or number + repeated_rows > MAX_ROWS:
                raise ValueError('В ODS слишком много строк. Разделите файл')
            for _ in range(repeated_rows):
                number += 1
                out.row(number, values)


def _docx(data, out):
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    relationship = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    with _zip(data) as archive:
        names = archive.namelist()
        root = _xml(archive.read('word/document.xml'))
        related = {}
        if 'word/_rels/document.xml.rels' in names:
            for rel in _xml(archive.read('word/_rels/document.xml.rels')):
                if rel.get('TargetMode') != 'External':
                    related[rel.get('Id')] = rel.get('Target', '')
        def text(element):
            parts = []
            for node in element.iter():
                if node.tag == w + 't':
                    parts.append(node.text or '')
                elif node.tag == w + 'tab':
                    parts.append('\t')
                elif node.tag in (w + 'br', w + 'cr'):
                    parts.append('\n')
            return ''.join(parts)
        def block(element):
            for node in element:
                if node.tag == w + 'p':
                    out.add(text(node))
                elif node.tag == w + 'tbl':
                    out.add('Таблица Word:')
                    for n, row in enumerate(node.findall(w + 'tr'), 1):
                        out.row(n, [text(c) for c in row.findall(w + 'tc')])
                else:
                    block(node)
        block(root)
        for name in sorted(names):
            if re.fullmatch(r'word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml', name):
                out.add('Раздел Word: ' + Path(name).stem)
                block(_xml(archive.read(name)))
        # Embedded price screenshots are source facts as well as body text.
        seen = set()
        for node in root.iter():
            rid = node.get(relationship + 'embed')
            target = related.get(rid)
            if not target or target in seen:
                continue
            seen.add(target)
            if '..' in Path(target).parts or target.startswith('/'):
                raise ValueError('Недопустимая ссылка на изображение Word')
            name = 'word/' + target
            if name in names:
                out.image(archive.read(name), 'Изображение Word')


def _pdf(data, out):
    import pymupdf
    if not data.lstrip().startswith(b'%PDF-'):
        raise ValueError('Файл PDF повреждён или имеет неверный формат')
    with pymupdf.open(stream=data, filetype='pdf') as doc:
        if doc.needs_pass:
            raise ValueError('PDF защищён паролем. Пришлите незашифрованную копию')
        if not 0 < doc.page_count <= MAX_PAGES:
            raise ValueError(f'PDF должен содержать от 1 до {MAX_PAGES} страниц. Разделите файл')
        for n, page in enumerate(doc, 1):
            out.add(f'Страница PDF {n}:')
            raw = page.get_text('text', sort=True)
            out.add(raw)
            area = max(1, page.rect.width * page.rect.height)
            image_area = sum(max(0, (v['bbox'][2] - v['bbox'][0]) * (v['bbox'][3] - v['bbox'][1]))
                             for v in page.get_image_info())
            if not raw.strip() or image_area / area > .15:
                if page.rect.width <= 0 or page.rect.height <= 0:
                    raise ValueError('В PDF повреждён размер страницы')
                scale = min(2.5, 2600 / max(page.rect.width, page.rect.height))
                pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                out.image(pix.tobytes('png'), f'Страница PDF {n}')


def _doc(data, out):
    converter = shutil.which('antiword')
    if not converter:
        raise ValueError('Для старого Word .doc нужен модуль antiword на сервере. Пока сохраните файл как .docx или PDF')
    with tempfile.TemporaryDirectory(prefix='supplier-doc-') as directory:
        source = Path(directory) / 'source.doc'
        source.write_bytes(data)
        result = subprocess.run([converter, '-m', 'UTF-8.txt', str(source)], capture_output=True,
                                timeout=30, check=False)
        if result.returncode:
            raise ValueError('Не удалось прочитать старый Word .doc. Сохраните файл как .docx или PDF')
        out.add(result.stdout.decode('utf-8'))


def extract(data: bytes, filename: str) -> dict:
    """Extract all accepted content or fail explicitly; never silently truncate.

    ``images`` contains bounded JPEGs for OCR/vision; ``text`` retains page,
    sheet and row labels. Original bytes should be retained by the caller.
    """
    if not isinstance(data, bytes) or not data:
        raise ValueError('Файл пуст или не содержит данных')
    if len(data) > MAX_FILE_BYTES:
        raise ValueError('Файл слишком большой — максимум 20 МБ')
    filename = str(filename).replace('\\', '/').split('/')[-1]
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError('Формат не поддерживается. Пришлите PDF, Word, TXT, CSV, TSV, YAML, JSON, Excel или ODS')
    out = _Output()
    try:
        if extension == '.txt':
            out.add(_decode(data))
        elif extension in ('.csv', '.tsv'):
            _csv(data, out, extension)
        elif extension in ('.json', '.yaml', '.yml'):
            _structured(data, out, extension)
        elif extension in ('.xlsx', '.xlsm'):
            _xlsx(data, out)
        elif extension == '.xls':
            _xls(data, out)
        elif extension == '.ods':
            _ods(data, out)
        elif extension == '.docx':
            _docx(data, out)
        elif extension == '.pdf':
            _pdf(data, out)
        elif extension == '.rtf':
            from striprtf.striprtf import rtf_to_text
            out.add(rtf_to_text(_decode(data), errors='strict'))
        else:
            _doc(data, out)
    except ImportError as exc:
        raise ValueError('На сервере не установлен модуль чтения этого формата. Требуется обновление зависимостей') from exc
    except (json.JSONDecodeError, UnicodeError) as exc:
        raise ValueError('Не удалось прочитать текст документа. Проверьте кодировку и синтаксис файла') from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Не удалось прочитать документ: файл повреждён, защищён или имеет неверный формат') from exc
    text = '\n'.join(out.lines)
    if not text.strip() and not out.images:
        raise ValueError('Документ пуст: текст или таблицы не найдены')
    return {'text': text, 'images': [{k: v for k, v in image.items() if k != '_size'} for image in out.images],
            'metadata': {'filename': filename, 'extension': extension, 'sha256': hashlib.sha256(data).hexdigest(),
                         'bytes': len(data), 'table_rows': out.rows, 'image_count': len(out.images),
                         'warnings': out.warnings}}


parse_document = extract
